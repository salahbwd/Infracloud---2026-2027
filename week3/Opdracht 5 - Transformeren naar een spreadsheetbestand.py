import openpyxl
from openpyxl.styles import Font, PatternFill

# 1. Bestandsnaam bepalen
excel_bestandsnaam = "systeeminformatie.xlsx"

# 2. Een nieuw Excel-werkboek en actief werkblad initialiseren
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Systeeminformatie"

# 3. Kolomnamen (headers) overnemen uit taak 4
kolomnamen = [
    "computernaam",
    "computermodel",
    "processor",
    "aantal_fysieke_cpu_cores",
    "maximale_cpu_frequentie_mhz",
    "totale_hoeveelheid_ram_gb",
    "bestandssysteem_opslag",
    "opslagcapaciteit_gb",
    "mac_adres_primaire_nic",
    "ipv4_adres_primaire_nic",
    "ipv6_adres_primaire_nic",
    "ipv6_link_local_adres_primaire_nic"
]

# Headers toevoegen als eerste rij
ws.append(kolomnamen)

# 4. Gegevensrij samenstellen met correcte datatypes
if 'resultaat' not in globals():
    resultaat = {}


rij_waarden = []
for veld in kolomnamen:
    waarde = resultaat.get(veld)
    rij_waarden.append("N/A" if waarde is None else waarde)


#Gegevens toevoegen als tweede rij
ws.append(rij_waarden)

# 5. Eenvoudige opmaak toevoegen (voordeel van een spreadsheetbestand)
header_font = Font(bold=True, color="FFFFFF")
header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")

for col_num in range(1, len(kolomnamen) + 1):
    cel = ws.cell(row=1, column=col_num)
    cel.font = header_font
    cel.fill = header_fill
    # Kolombreedte automatisch aanpassen aan de inhoud
    max_len = max(len(str(ws.cell(row=1, column=col_num).value or "")),
                  len(str(ws.cell(row=2, column=col_num).value or "")))
    col_letter = openpyxl.utils.get_column_letter(col_num)
    ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

# 6. Werkboek opslaan
wb.save(excel_bestandsnaam)
print(f"Gegevens succesvol geëxporteerd naar {excel_bestandsnaam}")
